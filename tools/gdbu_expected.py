"""Эталонный ГДБУ (Excel) → JSON для теста сверки: секции месторождений, бригады и их кусты по порядку.

    python tools/gdbu_expected.py "ГДБУ.xlsx" data/fixtures/gdbu.json [--sheet "Лист ГДБУ"]

Формат листа: заголовок секции «ГДБУ <месторождение> ГБ …», строки бригад (тип БУ, месторождение, № бригады),
дальше по месяцам ячейки «ВМР» и кампании «дд.мм.гггг / КП №N (скв) / дд.мм.гггг» — одной ячейкой
с переносами строк или тремя соседними ячейками.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from pathlib import Path

from openpyxl import load_workbook

DATE = re.compile(r"(\d{2})\.(\d{2})\.(\d{2,4})")


def tokens(ws, r: int, first_col: int) -> list[tuple[str, object]]:
    out = []
    for c in range(first_col, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v in (None, ""):
            continue
        if isinstance(v, dt.datetime):
            out.append(("d", v.date()))
            continue
        for part in str(v).split("\n"):
            part = part.strip()
            m = DATE.fullmatch(part)
            if m:
                y = int(m[3])
                out.append(("d", dt.date(y + 2000 if y < 100 else y, int(m[2]), int(m[1]))))
            elif part.upper() == "ВМР":
                out.append(("v", None))
            elif part:
                out.append(("l", part))
    return out


def campaigns(toks) -> list[dict]:
    out = []
    for i, (k, v) in enumerate(toks):
        if k != "l":
            continue
        m = re.search(r"\((\d+)\)\s*$", v)
        s = toks[i - 1][1] if i and toks[i - 1][0] == "d" else None
        e = toks[i + 1][1] if i + 1 < len(toks) and toks[i + 1][0] == "d" else None
        out.append({"pad": re.sub(r"\s*\(\d+\)\s*$", "", v).strip(), "wells": int(m[1]) if m else None,
                    "start": s.isoformat() if s else None, "end": e.isoformat() if e else None})
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("dst", type=Path)
    ap.add_argument("--sheet", default=None, help="лист с ГДБУ; по умолчанию — активный")
    a = ap.parse_args()
    wb = load_workbook(a.src, data_only=True)
    ws = wb[a.sheet] if a.sheet else wb.active
    sections, cur = [], None
    for r in range(1, ws.max_row + 1):
        title = str(ws.cell(r, 3).value or "")
        m = re.match(r"ГДБУ\s+(\S+)\s+ГБ", title)
        if m:
            cur = {"name": m[1], "rows": []}
            sections.append(cur)
            continue
        crew = ws.cell(r, 5).value
        if cur is None or crew in (None, "", "-") or str(crew).startswith("№"):
            continue
        cur["rows"].append({"crew": str(crew).strip().lstrip("0") or "0", "rig": str(ws.cell(r, 3).value or ""),
                            "campaigns": campaigns(tokens(ws, r, 6))})
    a.dst.parent.mkdir(parents=True, exist_ok=True)
    a.dst.write_text(json.dumps({"sections": sections}, ensure_ascii=False, indent=1), encoding="utf-8")
    for s in sections:
        print(s["name"], len(s["rows"]), "бригад,", sum(len(r["campaigns"]) for r in s["rows"]), "кампаний")


if __name__ == "__main__":
    main()
