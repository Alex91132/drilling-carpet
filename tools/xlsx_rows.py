"""Excel-график бурения → строки листа (как их видит страница) для тестов и демо-данных.

    python tools/xlsx_rows.py "ВПУ.xlsx" data/fixtures/vpu.json            # все столбцы — фикстура для тестов
    python tools/xlsx_rows.py "ВПУ.xlsx" "ТПУ.xlsx" --demo demo-data.js     # только нужные странице столбцы

В режиме --demo остаются только столбцы, которые читает парсер (куст, скважина, бригада, вид БУ, месторождение,
циклы и блоки фаз с датами). Подрядчик, инвестпроект, коды SAP, инвентарные и заводские номера не попадают.
Нужен openpyxl.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from pathlib import Path

from openpyxl import load_workbook

KEEP_BASE = {
    "порядок бурения скважины", "куст", "номер скважины", "проходка по скважине, м", "проходка по скважине",
    "цикл бурения", "цикл строительства", "коммерческая скорость, м/мес.", "коммерческая скорость, м/мес",
    "наименование вида бу", "номер бригады", "месторождение", "дата начала цикла бурения",
    "дата окончания цикла бурения", "количество грп/стадий мгрп", "количество гнкт",
    "классификация скважины по профилю ствола (короткое)", "классификация скважины по назначению (короткое)",
}
PHASE_GROUPS = {
    "бурение", "переезд", "освоение", "ввод скважины в эксплуатацию (вывод скважины на режим)", "ожидание освоения(испытания)",
    "ожидание освоения (испытания)", "консервация", "передвижка", "монтаж", "демонтаж", "обустройство (обвязка)",
    "стаскивание буровой установки", "бурение направления", "бурение кондуктора", "бурение технической колонны",
    "бурение эксплуатационной колонны", "бурение хвостовика",
}
PHASE_FIELDS = {"дата начала", "дата окончания", "количество дней", "расстояние на передвижку"}


def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def cell(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%dT%H:%M:%S") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, float):
        return round(v, 4)
    return v


def read_rows(path: Path) -> list[list]:
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    while rows and not any(v not in (None, "") for v in rows[-1]):
        rows.pop()
    return rows


def header_row(rows: list[list]) -> int:
    for i, r in enumerate(rows[:15]):
        names = {norm(v) for v in r}
        if "куст" in names and "номер скважины" in names:
            return i
    raise ValueError("нет строки заголовков с «Куст» и «Номер скважины»")


def demo_columns(rows: list[list]) -> list[int]:
    hr = header_row(rows)
    groups, cur = [], None
    for c in range(max(len(r) for r in rows[:hr + 1])):
        g = norm(rows[hr - 1][c]) if hr and c < len(rows[hr - 1]) else ""
        if g:
            cur = g if g in PHASE_GROUPS else None
        groups.append(cur)
    keep = []
    for c, h in enumerate(rows[hr]):
        h = norm(h)
        if h in KEEP_BASE or (h in PHASE_FIELDS and groups[c]):
            keep.append(c)
    return keep


def slim(v):
    """Для демо: дата без времени, два знака после запятой, нули — пусто (у фаз 0 дней = фазы нет)."""
    if isinstance(v, str) and re.match(r"\d{4}-\d{2}-\d{2}T", v):
        return v[:10]
    if isinstance(v, float):
        v = round(v, 2)
    return None if v == 0 else v


def trim(r: list) -> list:
    r = list(r)
    while r and r[-1] in (None, ""):
        r.pop()
    return r


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("--demo", type=Path, help="записать demo-data.js с нужными столбцами всех файлов")
    a = ap.parse_args()
    if a.demo:
        out = []
        for f in a.files:
            rows = read_rows(f)
            keep = demo_columns(rows)
            out.append({"name": re.sub(r"_.*$", "", f.stem),
                        "rows": [trim([slim(r[c]) if c < len(r) else None for c in keep]) for r in rows]})
            print(f"{f.name}: {len(rows)} строк, оставлено {len(keep)} столбцов из {len(rows[header_row(rows)])}")
        def js_rows(rows):   # пустые ячейки — «дыры» массива JS ([1,,"x"]): вдвое компактнее, чем null
            return "[" + ",\n".join("[" + ",".join("" if v is None else json.dumps(v, ensure_ascii=False) for v in r)
                                    + "]" for r in rows) + "]"
        body = ",\n".join("{name:" + json.dumps(f["name"], ensure_ascii=False) + ",rows:" + js_rows(f["rows"]) + "}"
                          for f in out)
        js = ("// Демо: графики бурения двух месторождений — только столбцы, которые читает страница.\n"
              "// Собрано tools/xlsx_rows.py --demo; формат тот же, что у загружаемого файла (строки листа).\n"
              "window.DEMO_FILES = [" + body + "];\n")
        a.demo.write_text(js, encoding="utf-8")
        print(f"{a.demo}: {len(js.encode()) / 1024:.0f} КБ")
        return
    src, dst = a.files[0], a.files[-1]
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(json.dumps({"name": src.stem, "rows": read_rows(src)}, ensure_ascii=False), encoding="utf-8")
    print(dst)


if __name__ == "__main__":
    main()
